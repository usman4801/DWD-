"""
DWD Tool
========
A Streamlit app that automates daily attendance ("DWD") report generation
for warehouse operations, styled as a polished Amazon Seller Tools page.
"""

import io
from datetime import datetime

import pandas as pd
import requests
import streamlit as st
from openpyxl import load_workbook

# =========================================================================
# CONFIGURATION
# =========================================================================

TEMPLATE_GITHUB_URL = "https://raw.githubusercontent.com/usman4801/DWD-/main/Blank%20file.xlsx"
GITHUB_TOKEN = st.secrets.get("GITHUB_TOKEN", "")

SHEET_ROSTER = "Roster"

# Exact raw CSV column mappings
COL_EMP_ID = "EmpID"
COL_DATE = "Date"
COL_PRESENT_STATUS = "present_status"
COL_IS_OVERTIME = "is_overtime"
COL_GB_LEAVE = "gb_leave_type"
COL_TIME_OFF = "time_off_type"

# Attendance codes based on template legend
CODE_PRESENT = "P"
CODE_PL = "PL"
CODE_OFF = "OFF"
CODE_ZERO = "0"

# Roster sheet layout based on Blank file.xlsx
ROSTER_HEADER_DATE_CELL = "B1"
ROSTER_HEADER_DAY_CELL = "C1"
ROSTER_ID_COLUMN = "B"
ROSTER_FIRST_DATA_ROW = 7
ROSTER_ATTENDANCE_COLUMN = "T"
ROSTER_REMARKS_COLUMN = "U"
ROSTER_OFF1_COLUMN = "L"
ROSTER_OFF2_COLUMN = "M"


# =========================================================================
# CORE LOGIC
# =========================================================================

@st.cache_data(show_spinner=False)
def fetch_template_bytes(url: str, token: str) -> bytes:
    """Download the blank template from GitHub as raw bytes."""
    headers = {"Accept": "application/vnd.github.v3.raw"}
    if token:
        headers["Authorization"] = f"token {token}"
    resp = requests.get(url, headers=headers, timeout=30)
    resp.raise_for_status()
    return resp.content


def build_report(raw_df: pd.DataFrame, template_bytes: bytes, selected_date) -> tuple:
    """
    Populate the Roster sheet of the template while keeping Dashboard formulas
    and all formatting/colors 100% intact.
    """
    day_name = selected_date.strftime("%A").strip().lower()

    wb = load_workbook(io.BytesIO(template_bytes), data_only=False)
    if SHEET_ROSTER not in wb.sheetnames:
        raise ValueError(f"Template is missing the '{SHEET_ROSTER}' sheet.")
    roster = wb[SHEET_ROSTER]

    roster[ROSTER_HEADER_DATE_CELL] = selected_date.strftime("%Y-%m-%d")
    roster[ROSTER_HEADER_DAY_CELL] = selected_date.strftime("%A")

    emp_data = {}
    for _, row in raw_df.iterrows():
        emp_id = str(row.get(COL_EMP_ID, "")).strip()
        if not emp_id or emp_id == "nan":
            continue
        present_status = str(row.get(COL_PRESENT_STATUS, "")).strip().lower()
        gb_leave = str(row.get(COL_GB_LEAVE, "")).strip().lower()
        time_off = str(row.get(COL_TIME_OFF, "")).strip().lower()
        is_ot = row.get(COL_IS_OVERTIME, 0)

        is_present = present_status == "present"
        is_pl = "annual leave" in gb_leave or "annualleave" in time_off

        emp_data[emp_id] = {
            "is_present": is_present,
            "is_pl": is_pl,
            "is_ot": is_ot == 1 or is_ot == "1",
        }

    row_idx = ROSTER_FIRST_DATA_ROW
    while True:
        id_cell = roster[f"{ROSTER_ID_COLUMN}{row_idx}"]
        if id_cell.value in (None, ""):
            break
        emp_id = str(id_cell.value).strip()

        data = emp_data.get(emp_id)
        if data:
            off1_cell = roster[f"{ROSTER_OFF1_COLUMN}{row_idx}"]
            off2_cell = roster[f"{ROSTER_OFF2_COLUMN}{row_idx}"]

            off1_val = str(off1_cell.value or "").strip().lower()
            off2_val = str(off2_cell.value or "").strip().lower()

            is_scheduled_off = (day_name == off1_val or day_name == off2_val)

            if data["is_pl"]:
                roster[f"{ROSTER_ATTENDANCE_COLUMN}{row_idx}"] = CODE_PL
                roster[f"{ROSTER_REMARKS_COLUMN}{row_idx}"] = "Annual Leave"
            elif data["is_present"]:
                roster[f"{ROSTER_ATTENDANCE_COLUMN}{row_idx}"] = CODE_PRESENT

                if day_name == off1_val:
                    off1_cell.value = "OT"
                    roster[f"{ROSTER_REMARKS_COLUMN}{row_idx}"] = "6th day OT"
                elif day_name == off2_val:
                    off2_cell.value = "OT"
                    roster[f"{ROSTER_REMARKS_COLUMN}{row_idx}"] = "7th day OT"
            elif is_scheduled_off:
                roster[f"{ROSTER_ATTENDANCE_COLUMN}{row_idx}"] = CODE_OFF
            else:
                roster[f"{ROSTER_ATTENDANCE_COLUMN}{row_idx}"] = CODE_ZERO

        row_idx += 1

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out.getvalue(), selected_date


# =========================================================================
# UI — Amazon Seller Tools theme, polished, single viewport
# =========================================================================

st.set_page_config(page_title="DWD Tool - Amazon Operations", page_icon="📦", layout="wide")

st.markdown(
    """
    <style>
    #MainMenu, header, footer {visibility: hidden;}
    html, body {height: 100%; margin: 0; padding: 0;}
    .stApp {background: #eef1f5; min-height: 100vh;}
    .block-container {padding: 0 !important; max-width: 100% !important;}

    .amz-navbar {
        background: #0e1420;
        padding: 9px 40px;
        display: flex;
        justify-content: space-between;
        align-items: center;
        border-bottom: 1px solid #232f3e;
    }
    .amz-logo {
        color: #ffffff;
        font-size: 1.3rem;
        font-weight: 800;
        font-style: italic;
    }
    .amz-logo::after {
        content: "⌣";
        color: #ff9900;
        font-size: 1.05rem;
        margin-left: 2px;
    }
    .amz-nav-icons {
        display: flex;
        align-items: center;
        gap: 18px;
        color: #d9dee6;
    }
    .amz-nav-icons .avatar {
        width: 27px; height: 27px;
        border-radius: 50%;
        background: linear-gradient(135deg,#ff9900,#7c4fe0);
        display: flex; align-items: center; justify-content: center;
        font-size: 0.8rem;
        color: #fff;
    }

    .amz-hero {
        background: linear-gradient(120deg, #0b1524 0%, #14273e 45%, #223f5f 100%);
        padding: 14px 40px 12px 40px;
    }
    .hero-flex {display: flex; justify-content: space-between; align-items: center;}
    .amz-badge {
        display: inline-flex;
        align-items: center;
        gap: 6px;
        background: rgba(255,153,0,0.16);
        border: 1px solid rgba(255,153,0,0.55);
        color: #ffb54d;
        font-weight: 600;
        font-size: 0.7rem;
        padding: 3px 12px;
        border-radius: 20px;
        margin-bottom: 8px;
    }
    .amz-hero h1 {
        color: #ffffff;
        font-size: 2.15rem;
        font-weight: 800;
        margin: 0 0 2px 0;
    }
    .amz-hero h1 span {
        background: linear-gradient(90deg, #ffb230, #ff8a00);
        -webkit-background-clip: text;
        background-clip: text;
        color: transparent;
    }
    .amz-hero h2 {
        color: #eef2f7;
        font-size: 1.02rem;
        font-weight: 600;
        margin: 0 0 6px 0;
    }
    .amz-hero p.desc {
        color: #aebdd1;
        font-size: 0.82rem;
        max-width: 460px;
        margin-bottom: 12px;
    }

    .amz-content {padding: 10px 40px 4px 40px;}
    .amz-panel {
        background: #ffffff;
        border-radius: 14px;
        box-shadow: 0 10px 26px rgba(20,30,50,0.1);
        border: 1px solid #eef0f4;
        padding: 16px;
    }
    .left-card {
        background: linear-gradient(160deg,#f5f8fd,#eef2fa);
        border-radius: 12px;
        padding: 18px 16px;
        height: 100%;
        border: 1px solid #e6ecf6;
    }
    .left-card h3 {
        color: #131921;
        font-size: 1.05rem;
        font-weight: 800;
        margin: 0 0 8px 0;
    }
    .left-card p {
        color: #5b6673;
        font-size: 0.79rem;
        margin-bottom: 10px;
    }

    div.stButton>button {
        background: linear-gradient(135deg,#ffb230,#ff8a00);
        color: #14202e;
        font-weight: 700;
        border-radius: 8px;
        padding: 0.36rem 1.35rem;
        border: none;
    }
    div.stDownloadButton>button {
        background: linear-gradient(135deg,#33cb95,#0f9d68);
        color: #ffffff;
        font-weight: 700;
        border-radius: 8px;
        padding: 0.36rem 1.35rem;
        border: none;
    }
    </style>
    """,
    unsafe_allow_html=True,
)

# =========================================================================
# ACCESS GATE — Simple login lock screen
# =========================================================================
if "unlocked" not in st.session_state:
    st.session_state.unlocked = False

if not st.session_state.unlocked:
    col1, col2, col3 = st.columns([1, 1.2, 1])
    with col2:
        st.markdown("<br><br>", unsafe_allow_html=True)
        st.markdown(
            """
            <div style="background: white; padding: 30px; border-radius: 16px; box-shadow: 0 15px 35px rgba(0,0,0,0.2); text-align: center; border-top: 5px solid #ff9900;">
                <div style="font-size: 2rem; margin-bottom: 10px;">🔒</div>
                <h3 style="color: #131921; margin-bottom: 5px;">DWD Tool Login</h3>
                <p style="color: #6b7280; font-size: 0.85rem; margin-bottom: 20px;">Please enter your access code to unlock</p>
            </div>
            """,
            unsafe_allow_html=True,
        )
        entered_code = st.text_input("Access Code", type="password", label_visibility="collapsed", placeholder="Enter code here")
        if entered_code:
            if entered_code.strip().lower() == "javmuhak":
                st.session_state.unlocked = True
                st.rerun()
            else:
                st.error("Incorrect code. Please try again.")
        st.stop()


# --- Navbar ---
st.markdown(
    """
    <div class="amz-navbar">
        <div class="amz-logo">amazon</div>
        <div class="amz-nav-icons">
            <span>⠿</span>
            <span>🔔</span>
            <div class="avatar">👤</div>
        </div>
    </div>
    """,
    unsafe_allow_html=True,
)

# --- Hero ---
st.markdown(
    """
    <div class="amz-hero">
        <div class="hero-flex">
            <div>
                <div class="amz-badge">⚡ Amazon Tool</div>
                <h1>DWD <span>Tool</span></h1>
                <h2>Amazon Daily Workforce Dashboard Generator</h2>
                <p class="desc">
                    Generate your daily workforce dashboard report quickly and easily.
                    Upload your data and get actionable insights in just a few clicks.
                </p>
            </div>
        </div>
    </div>
    """,
    unsafe_allow_html=True,
)

# --- Main content ---
st.markdown('<div class="amz-content">', unsafe_allow_html=True)

left_col, right_col = st.columns([1, 2.1], gap="medium")

with left_col:
    st.markdown(
        """
        <div class="left-card">
            <h3>Create Your Dashboard</h3>
            <p>
                Select the report date and upload your raw data file in CSV
                format to generate your Amazon workforce dashboard.
            </p>
        </div>
        """,
        unsafe_allow_html=True,
    )

with right_col:
    st.markdown('<div class="amz-panel">', unsafe_allow_html=True)
    st.markdown("#### 1. Select Report Date")
    selected_date = st.date_input("Select Report Date", value=datetime.today(), label_visibility="collapsed")

    st.markdown("#### 2. Upload Raw File (CSV)")
    uploaded_file = st.file_uploader("Upload Raw File (CSV)", type=["csv"], label_visibility="collapsed")

    generate_clicked = st.button("📊  Download DWD  →")

    if generate_clicked:
        if uploaded_file is None:
            st.error("Please upload a raw CSV file first.")
        else:
            try:
                raw_df = pd.read_csv(uploaded_file)
            except Exception as e:
                st.error(f"Error reading raw CSV file: {e}")
                raw_df = None

            if raw_df is not None:
                with st.spinner("Generating DWD report..."):
                    try:
                        template_bytes = fetch_template_bytes(TEMPLATE_GITHUB_URL, GITHUB_TOKEN)
                        report_bytes, report_date = build_report(raw_df, template_bytes, selected_date)
                        st.session_state["report_bytes"] = report_bytes
                        st.session_state["report_date"] = report_date
                    except Exception as e:
                        st.error(f"Failed to generate report: {e}")

    if "report_bytes" in st.session_state:
        st.download_button(
            "⬇  Save DWD Report",
            data=st.session_state["report_bytes"],
            file_name=f"DWD-AUH1-{st.session_state['report_date'].strftime('%d%m%Y')}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    st.markdown('</div>', unsafe_allow_html=True)

st.markdown('</div>', unsafe_allow_html=True)
